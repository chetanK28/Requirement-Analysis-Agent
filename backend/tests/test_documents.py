from io import BytesIO
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest
from fastapi.testclient import TestClient
from starlette.datastructures import UploadFile, Headers
from app.main import app
from app.services.documents import MAX_FILE_BYTES, DOCX_TYPE
from app.api.routes import extract_file
from app.providers.base import ProviderError


def pdf(text='As a customer I want to reset my password', encrypted=False, blank=False):
    writer=PdfWriter();page=writer.add_blank_page(width=612,height=792)
    if not blank:
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
    if encrypted:writer.encrypt('test-password')
    output=BytesIO();writer.write(output);return output.getvalue()


def docx(empty=False,text='Original paragraph'):
    doc=Document()
    if not empty:
        doc.add_paragraph(text)
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Table requirement';table.cell(0,1).text='Table acceptance criteria'
        doc.add_paragraph('Final paragraph')
    output=BytesIO();doc.save(output);return output.getvalue()


@pytest.fixture
def client():
    with TestClient(app) as client:yield client


def send(client,data,name='story.pdf',mime='application/pdf'):
    return client.post('/api/v1/extract-file',files={'file':(name,data,mime)})


def test_valid_pdf_does_not_call_model(client):
    with patch('app.providers.openai_provider.AsyncOpenAI') as model:
        response=send(client,pdf())
        model.assert_not_called()
    assert response.status_code==200
    data=response.json();assert 'reset my password' in data['text'];assert data['character_count']==len(data['text'])


def test_docx_paragraphs_and_table_preserve_order(client):
    response=send(client,docx(),'STORY.DOCX',DOCX_TYPE)
    assert response.status_code==200
    assert response.json()['text'].splitlines()==['Original paragraph','Table requirement','Table acceptance criteria','Final paragraph']


@pytest.mark.parametrize('name,mime',[('old.doc','application/msword'),('bad.txt','text/plain'),('bad.exe','application/octet-stream')])
def test_invalid_extensions(client,name,mime):
    response=send(client,b'private source',name,mime)
    assert response.status_code==415
    assert 'private source' not in response.text
    if name.endswith('.doc'):assert 'Legacy .doc' in response.json()['error']['message']


@pytest.mark.parametrize('name,mime',[('fake.pdf',DOCX_TYPE),('fake.docx','application/pdf'),('story.pdf','application/octet-stream')])
def test_content_type_mismatch(client,name,mime):
    assert send(client,b'content',name,mime).status_code==415


def test_empty_upload(client):
    response=send(client,b'');assert response.status_code==422;assert response.json()['error']['code']=='empty_file'


def test_oversized_upload(client):
    response=send(client,b'x'*(MAX_FILE_BYTES+1));assert response.status_code==413


@pytest.mark.parametrize('name,mime',[('bad.pdf','application/pdf'),('bad.docx',DOCX_TYPE)])
def test_corrupt_document(client,name,mime):
    response=send(client,b'private corrupt document',name,mime)
    assert response.status_code==422
    assert response.json()['error']['code']=='corrupt_document'
    assert 'private corrupt document' not in response.text


def test_scan_or_blank_pdf_explains_no_ocr(client):
    response=send(client,pdf(blank=True));assert response.status_code==422
    assert 'OCR' in response.json()['error']['message']
    assert 'paste' in response.json()['error']['message']


def test_password_protected_pdf(client):
    response=send(client,pdf(encrypted=True));assert response.status_code==422
    assert response.json()['error']['code']=='password_protected'


def test_empty_docx(client):
    response=send(client,docx(empty=True),'empty.docx',DOCX_TYPE)
    assert response.status_code==422;assert response.json()['error']['code']=='no_extractable_text'


def test_extracted_text_limit_no_silent_truncation(client):
    response=send(client,docx(text='x'*12001),'long.docx',DOCX_TYPE)
    assert response.status_code==422
    assert response.json()['error']['code']=='text_too_long'
    assert 'No text was truncated' in response.json()['error']['message']


def test_upload_handle_closed_even_on_failure():
    import asyncio
    file=UploadFile(BytesIO(b'broken'),filename='broken.pdf',headers=Headers({'content-type':'application/pdf'}))
    with pytest.raises(ProviderError):asyncio.run(extract_file(file))
    assert file.file.closed

def test_oversized_body_rejected_before_parser_and_extraction(client):
    with patch('app.services.documents.extract_text') as extract:
        response=send(client,b'x'*(MAX_FILE_BYTES+65537))
        extract.assert_not_called()
    assert response.status_code==413


def test_missing_file_has_upload_guidance(client):
    response=client.post('/api/v1/extract-file',data={})
    assert response.status_code==422
    assert response.json()['error']['code']=='invalid_upload'


def test_nested_docx_table_text(client):
    doc=Document();cell=doc.add_table(rows=1,cols=1).cell(0,0)
    cell.text='Outer cell';cell.add_table(rows=1,cols=1).cell(0,0).text='Nested requirement'
    output=BytesIO();doc.save(output)
    response=send(client,output.getvalue(),'nested.docx',DOCX_TYPE)
    assert response.status_code==200
    assert 'Outer cell' in response.json()['text'] and 'Nested requirement' in response.json()['text']